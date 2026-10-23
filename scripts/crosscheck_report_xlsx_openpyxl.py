#!/usr/bin/env python3
"""Read workbench.reports.write_xlsx output back with openpyxl 3.1.5 and compare every cell.

openpyxl is not a Polymer Lab dependency. Install it beside the project environment, for example:
  .venv/bin/python -m pip install --target /tmp/xlsxlib --no-deps openpyxl==3.1.5 et_xmlfile==2.0.0
  PYTHONPATH=/tmp/xlsxlib .venv/bin/python scripts/crosscheck_report_xlsx_openpyxl.py . [reports folder ...]
"""
import csv
import math
from pathlib import Path
import random
import sys
import tempfile

from openpyxl import load_workbook

sys.path.insert(0, sys.argv[1] if len(sys.argv) > 1 else ".")
from workbench import reports  # noqa: E402


def expected(value):
    value = reports._cell(value)
    if isinstance(value, str):
        return reports._XML_INVALID.sub("", value)[:32767] or None
    if isinstance(value, float) and not math.isfinite(value):
        return str(value)
    return value


def same(a, b):
    if isinstance(a, float) or isinstance(b, float):
        return isinstance(a, (int, float)) and isinstance(b, (int, float)) and not isinstance(a, bool) and float(a) == float(b)
    return a == b and type(a) is type(b)


rng = random.Random(2026)
pool = ["온도 <K> & \"압력\"", "  leading and trailing  ", "=SUM(A1:A2)", "tab\tand\nnewline", "ctrl\x01\x07chars", "é ü 汉字 한글 😀",
        "", "x" * 40000, "</t></is></c>", "]]>"]
sheets = []
for index in range(6):
    width = rng.randint(1, 40)
    columns = [f"col {j} {rng.choice(pool[:6])}" for j in range(width)]
    rows = []
    for _ in range(rng.randint(0, 60)):
        row = []
        for _ in range(width):
            kind = rng.random()
            row.append(None if kind < 0.1 else rng.choice([True, False]) if kind < 0.2 else rng.randint(-10**12, 10**12) if kind < 0.35
                       else rng.choice([0.0, -0.0, 5e-324, 1.7976931348623157e308, rng.uniform(-1e6, 1e6), rng.lognormvariate(0, 20), math.inf])
                       if kind < 0.7 else rng.choice(pool) if kind < 0.9 else {"nested": [1, 2.5, "한글"]})
        rows.append(row)
    sheets.append((["Sheet", "Sheet", "bad[]:*?/\\name", "a" * 50, "표", "Sheet"][index], columns, rows))
cells = 0
with tempfile.TemporaryDirectory() as temp:
    path = Path(temp) / "check.xlsx"
    reports.write_xlsx(path, sheets)
    book = load_workbook(path)
    if len(book.worksheets) != len(sheets) or len({s.title.lower() for s in book.worksheets}) != len(sheets):
        raise SystemExit("sheet names are not unique or a sheet is missing")
    for sheet, (_, columns, rows) in zip(book.worksheets, sheets):
        if sheet.freeze_panes != "A2":
            raise SystemExit(f"{sheet.title}: header row is not frozen")
        grid = [list(r) for r in sheet.iter_rows(values_only=True)]
        for r, row in enumerate([columns, *rows]):
            for c, value in enumerate(row):
                got = grid[r][c] if r < len(grid) and c < len(grid[r]) else None
                if not same(expected(value), got) and not (expected(value) is None and got is None):
                    raise SystemExit(f"{sheet.title} row {r + 1} col {c + 1}: wrote {expected(value)!r}, openpyxl read {got!r}")
                cells += 1
print(f"{len(sheets)} synthetic sheets, {cells} cells (Unicode, XML specials, control characters, 40 000-character text, extreme floats, bools, JSON): identical")

reports_checked = 0
for folder in sys.argv[2:]:
    for workbook in sorted(Path(folder).glob("*/content/tables/report.xlsx")):
        book = load_workbook(workbook)
        tables = sorted(p for p in workbook.parent.glob("*.csv"))
        by_header = {}
        for table in tables:
            with table.open(encoding="utf-8-sig", newline="") as stream:
                rows = list(csv.reader(stream))
            by_header.setdefault(tuple(rows[0]), []).append(rows)
        for sheet in book.worksheets:
            grid = [["" if v is None else ("TRUE" if v is True else "FALSE" if v is False else repr(v) if isinstance(v, float) else str(v)) for v in r]
                    for r in sheet.iter_rows(values_only=True)]
            candidates = by_header.get(tuple(grid[0]), [])
            def matches(rows):
                return len(rows) == len(grid) and all(
                    len(a) == len(b) and all(x == y or x == "'" + y or (x and y and _float_equal(x, y)) for x, y in zip(a, b)) for a, b in zip(rows, grid))
            def _float_equal(x, y):
                try:
                    return float(x) == float(y)
                except ValueError:
                    return False
            if not any(matches(rows) for rows in candidates):
                raise SystemExit(f"{workbook}: sheet {sheet.title} differs from its CSV table")
        reports_checked += 1
if len(sys.argv) > 2:
    print(f"{reports_checked} report workbooks: every sheet equals its CSV table")
